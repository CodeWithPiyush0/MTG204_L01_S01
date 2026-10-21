# -*- coding: utf-8 -*-
"""Build the VO manifest for the SME — MTG2A04_L01_S01_VO_MANIFEST.xlsx in the game folder.

Same layout as the reference lesson's lesson_vo_manifest.xlsx (HI02H11_L02_S02): four sheets —
VO Manifest (every clip in game order, where and when it plays), Recording list (each recording once),
Sound effects (supplied, not recorded), Recording spec. The Hindi script is read from card.json (the
build's own spoken text), durations from the clips on disk, so re-running this after any VO change
keeps it exact.

Run (from the game folder):   PYTHONUTF8=1 python scripts/gen_vo_manifest.py
"""
import os, json, subprocess, datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), ".."))
AUD = os.path.join(ROOT, "assets", "Audio")
OUT = os.path.join(ROOT, "MTG2A04_L01_S01_VO_MANIFEST.xlsx")
CARD = json.load(open(os.path.join(ROOT, "card.json"), encoding="utf-8"))
TEXT = dict(CARD["assets"]["audio_text"])
# the guided / practice transition lines are the reference lesson's clips (not registered by this build)
TEXT.setdefault("vo_pt_guided", "बहुत बढ़िया! अब हम साथ मिलकर शुरू करते हैं। चलिए, साथ में करें!")
TEXT.setdefault("vo_pt_practice", "वाह! अब आपकी बारी।")

DRAFT = {"vo_t1_note", "vo_t1_coin"}                 # drafted by the developer, the SME to finalise
EAR = {"vo_b45_prompt"}                              # TTS needed a retry (punctuation variant)
SRC = os.path.join(AUD, "voiceovers")                # the SME's recordings (prepared by scripts/prepare_vo.py)
# where the recording's words differ from the script (checked by transcription, 2026-10-21)
SAYS = {}   # (round 16: vo_pt_tutorial's script now says «…चलिए, शुरू करें!», as recorded)

def dur(i):
    p = os.path.join(AUD, i + ".ogg")
    if not os.path.isfile(p): return None
    out = subprocess.run(["ffprobe", "-v", "error", "-show_entries", "format=duration", "-of", "csv=p=0", p],
                         capture_output=True, text=True).stdout.strip()
    return round(float(out), 2) if out else None

# (page, screen, audio id, what it is, when it plays, same recording also on)
P1, P2, P3, P4, P5, P6, P7, P8 = ("Page %d" % n for n in range(1, 9))
T1, T2, T3 = "Transition 1 (before page 1)", "Transition 2 (before page 2)", "Transition 3 (before page 7)"
ROWS = [
 ("Start screen", "Start screen", "vo_landing", "greeting", "Start screen - Swiftee's greeting as it loads (the speaker button replays it); the play button appears when it ends", ""),
 (T1, "Transition «चलिए, शुरू करें!»", "vo_pt_tutorial", "transition line", "After the play button - Swiftee rises, then says the line; the title types in as she says it", ""),
 (P1, "Teaching - ₹10 and ₹1", "vo_t1_note", "introduces the ₹10 note", "Page opens - the ₹10 note and its «₹10» label appear WITH this line", ""),
 (P1, "Teaching - ₹10 and ₹1", "vo_t1_coin", "introduces the ₹1 coin", "Right after - the ₹1 coin and its «₹1» label appear WITH this line; then the arrow", ""),
 (T2, "Transition «चलिए, साथ में करें।»", "vo_pt_guided", "transition line", "After page 1 - Swiftee rises, then says the line; the title types in as she says it", ""),
 (P2, "Sort ₹10 (notes + coins into the purse)", "vo_g1_prompt", "page instruction", "After the notes and coins have popped in - the instruction text shows with it; then the hand taps a ₹10 coin", ""),
 (P2, "Sort ₹10 (notes + coins into the purse)", "vo_sort_done", "praise", "All four ₹10 pieces are in the purse (with confetti); then the arrow", "Page 3; Page 7 (सेब, टमाटर, अंगूर)"),
 (P3, "Sort ₹1 (coins into the purse)", "vo_g2_prompt", "page instruction", "After the coins have popped in - the instruction text shows with it", ""),
 (P3, "Sort ₹1 (coins into the purse)", "vo_sort_done", "praise (plays again)", "All four ₹1 coins are in the purse (with confetti); then the arrow", "Page 2; Page 7"),
 (P4, "Make ₹14 (ghost-guided)", "vo_b14_intro", "opening line", "Page opens - Swiftee's line (the panel shows it)", ""),
 (P4, "Make ₹14 (ghost-guided)", "vo_b14_prompt", "page instruction", "Right after - then a ghost shows each move: one ₹10, then each ₹1", ""),
 (P4, "Make ₹14 (ghost-guided)", "vo_g2_done", "closing line", "₹14 made - the ₹10 row, the ₹1 row and the total light up as the line names them; then the arrow", ""),
 (P5, "Make ₹32 (₹10 first, then ₹1)", "vo_b32_prompt", "page instruction", "Page opens - ₹10 open, ₹1 locked", ""),
 (P5, "Make ₹32 (₹10 first, then ₹1)", "vo_hint_ones_2", "Hint 2 (2nd extra ₹10)", "The child tries a 4th / 5th ₹10 the 2nd time: it bounces back + a ghost ₹1 + this line (₹2 still missing). 1st extra ₹10: bounce + shake, no voice; 3rd: the ₹10 locks", "Page 6"),
 (P5, "Make ₹32 (₹10 first, then ₹1)", "vo_hint_ones_1", "Hint 2 - variant", "Same moment, but one ₹1 is already placed (₹1 still missing)", "Page 6"),
 (P5, "Make ₹32 (₹10 first, then ₹1)", "vo_b32_done", "closing line", "₹32 made - rows and total light up as named; then the arrow", ""),
 (P6, "Make ₹45 (₹10 first, then ₹1)", "vo_b45_prompt", "page instruction", "Page opens - ₹10 open, ₹1 locked", ""),
 (P6, "Make ₹45 (₹10 first, then ₹1)", "vo_hint_ones_5", "Hint 2 (2nd extra ₹10)", "The child tries a 5th ₹10 the 2nd time: it bounces back + a ghost ₹1 + this line (₹5 still missing) - the usual case", ""),
 (P6, "Make ₹45 (₹10 first, then ₹1)", "vo_hint_ones_4", "Hint 2 - variant", "Same moment, one ₹1 already placed (₹4 missing)", ""),
 (P6, "Make ₹45 (₹10 first, then ₹1)", "vo_hint_ones_3", "Hint 2 - variant", "Same moment, two ₹1 already placed (₹3 missing)", ""),
 (P6, "Make ₹45 (₹10 first, then ₹1)", "vo_hint_ones_2", "Hint 2 - variant (plays again)", "Same moment, three ₹1 already placed (₹2 missing)", "Page 5"),
 (P6, "Make ₹45 (₹10 first, then ₹1)", "vo_hint_ones_1", "Hint 2 - variant (plays again)", "Same moment, four ₹1 already placed (₹1 missing)", "Page 5"),
 (P6, "Make ₹45 (₹10 first, then ₹1)", "vo_b45_done", "closing line", "₹45 made - rows and total light up as named; then the arrow", ""),
 (T3, "Transition «अब आपकी बारी।»", "vo_pt_practice", "transition line", "After page 6 - Swiftee rises, then says the line; the title types in as she says it", ""),
 (P7, "Shop game - the stall", "vo_m_intro", "opening line", "The stall opens - Swiftee introduces the game (the panel shows the line)", ""),
 (P7, "Shop game - the stall", "vo_m_prompt", "page instruction", "Right after; the speaker button replays it", ""),
 (P7, "Shop game - the stall", "vo_m_idle", "nudge", "On the stall, 7 s without a tap - an item glows with this line (up to 3 times per visit)", ""),
 (P7, "Shop game - make the amount", "vo_m_sel_apple", "item line - सेब ₹14", "सेब tapped - its make-the-amount screen opens", ""),
 (P7, "Shop game - make the amount", "vo_m_sel_oranges", "item line - संतरे ₹25", "संतरे tapped", ""),
 (P7, "Shop game - make the amount", "vo_m_sel_tomatoes", "item line - टमाटर ₹35", "टमाटर tapped", ""),
 (P7, "Shop game - make the amount", "vo_m_sel_grapes", "item line - अंगूर ₹42", "अंगूर tapped", ""),
 (P7, "Shop game - make the amount", "vo_m_sel_bananas", "item line - केले ₹60", "केले tapped", ""),
 (P7, "Shop game - make the amount", "vo_m_sel_carrots", "item line - गाजर ₹65", "गाजर tapped", ""),
 (P7, "Shop game - make the amount", "vo_m_tip", "one-time tip", "The first money placed in the shop (once per game) - a ghost shows how to drag money back out", ""),
 (P7, "Shop game - make the amount", "vo_m_less1", "Hint 1 - too little", "1st wrong «जाँचें» with too little money - the money stays", ""),
 (P7, "Shop game - make the amount", "vo_m_more1", "Hint 1 - too much", "1st wrong «जाँचें» with too much money - the money stays", ""),
 (P7, "Shop game - make the amount", "vo_m_less2", "Hint 2 - still too little", "2nd wrong «जाँचें» with too little - the money stays", ""),
 (P7, "Shop game - make the amount", "vo_m_more2", "Hint 2 - still too much", "2nd wrong «जाँचें» with too much - the money stays", ""),
 (P7, "Shop game - make the amount", "vo_m_retry", "Hint 3 - start again", "3rd wrong «जाँचें» - the tray empties (every piece bounces home); the count starts again", ""),
 (P7, "Shop game - make the amount", "vo_m_ok_oranges", "praise «बहुत बढ़िया!»", "Correct amount - the FIRST purchase of the game (any item), and every later purchase of संतरे, केले, गाजर", ""),
 (P7, "Shop game - make the amount", "vo_sort_done", "praise «शाबाश!» (plays again)", "Correct amount - later purchases of सेब, टमाटर, अंगूर", "Pages 2, 3"),
 (P7, "Shop game - all bought", "vo_m_done", "closing line", "All six bought - Pari with her full basket and the ticked list", ""),
 (P8, "Celebration", "vo_cel_prompt", "end-screen dialogue", "After Swiftee's jump - she says the line, lip-synced; then the arrow pulses", ""),
]
missing = sorted(set(TEXT) - {r[2] for r in ROWS})
assert not missing, "clips not placed in the manifest: %s" % missing

SFX = [
 ("sfx_play_button", "Play button on the start screen - sounds on the press", "Start"),
 ("sfx_next_button", "Next arrow, and the arrow on the celebration screen (real presses only)", "all pages"),
 ("sfx_drop_note", "A note lands - in the drop zone or the purse", "2, 4-7"),
 ("sfx_drop_coin", "A coin lands - in the drop zone or the purse", "2-7"),
 ("sfx_fb_correct", "Correct - a screen completed, a purchase made", "2-7"),
 ("sfx_fb_incorrect", "Incorrect - a wrong note / coin, an extra ₹10, a wrong «जाँचें»", "2-7"),
 ("sfx_confetti", "With every confetti burst (team standard pack)", "2-7"),
 ("sfx_nudge", "Soft chime with every 7 s inactivity nudge (team standard pack)", "2-7"),
 ("sfx_celebrate", "Celebration screen", "8"),
 ("bgm_lesson", "Background music - from the play button on, ducked under voice and sounds", "all"),
]

NAVY, GREY = "1F4E79", "555555"
F_TITLE = Font(bold=True, size=14, color=NAVY); F_SUB = Font(size=10, color=GREY)
F_HEAD = Font(bold=True, color="FFFFFF"); FILL_HEAD = PatternFill("solid", fgColor=NAVY)
F_HI = Font(name="Nirmala UI", size=12)
FILL_TR, FILL_A, FILL_B = (PatternFill("solid", fgColor=c) for c in ("EEE8FB", "FFFFFF", "F3F7FB"))
THIN = Side(style="thin", color="D9D9D9"); BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="center")

def sheet(ws, title, sub, head, widths, freeze="E5"):
    ws["A1"] = title; ws["A1"].font = F_TITLE
    ws["A2"] = sub; ws["A2"].font = F_SUB
    for j, h in enumerate(head, 1):
        c = ws.cell(4, j, h); c.font = F_HEAD; c.fill = FILL_HEAD; c.alignment = WRAP; c.border = BOX
    for col, w in widths.items(): ws.column_dimensions[col].width = w
    if freeze: ws.freeze_panes = freeze

def status(i):
    if not os.path.isfile(os.path.join(AUD, i + ".ogg")): return "Missing - to record"
    if os.path.isfile(os.path.join(SRC, i + ".wav")):
        s = "Recorded (SME, %s)" % datetime.date.fromtimestamp(os.path.getmtime(os.path.join(SRC, i + ".wav"))).isoformat()
        if i in SAYS: s += " - says «%s»" % SAYS[i].split("। ")[-1]
        return s
    s = "TTS placeholder - to record"
    if i in DRAFT: s += " - DRAFT script, SME to finalise"
    if i in EAR: s += " - EAR-CHECK"
    return s

wb = Workbook()
# ---- VO Manifest
ws = wb.active; ws.title = "VO Manifest"
distinct = []
for r in ROWS:
    if r[2] not in distinct: distinct.append(r[2])
sheet(ws, "परी चली बाज़ार (MTG2A04_L01_S01) - lesson voice-over manifest",
      "Start screen -> transitions and pages 1-8 in game order; on each page the clips in the order they play. "
      "%d rows, %d distinct recordings. Folder assets/Audio/  ·  current takes are Gemini TTS placeholders (voice Leda)" % (len(ROWS), len(distinct)),
      ["S.No", "Page", "Screen", "#", "Audio ID", "Hindi script (spoken exactly)", "What it is", "When it plays",
       "Same recording also on", "Duration (s)", "Status"],
      {"A": 6, "B": 15, "C": 26, "D": 4, "E": 20, "F": 48, "G": 26, "H": 62, "I": 20, "J": 10, "K": 30})
band, last_page, n_on_page = 0, None, 0
for k, (page, screen, aid, what, when, also) in enumerate(ROWS, 1):
    if page != last_page:
        n_on_page = 0
        if page.startswith(("Page",)): band += 1
        last_page = page
    n_on_page += 1
    fill = FILL_TR if not page.startswith("Page") else (FILL_A if band % 2 else FILL_B)
    vals = [k, page, screen, n_on_page, aid, TEXT[aid], what, when, also or None, dur(aid), status(aid)]
    for j, v in enumerate(vals, 1):
        c = ws.cell(4 + k, j, v); c.fill = fill; c.alignment = WRAP; c.border = BOX
        if j == 6: c.font = F_HI
# ---- Recording list
ws = wb.create_sheet("Recording list")
sheet(ws, "Recording list - each recording once, in the order it is first heard",
      "%d recordings. Record each once; the file name is the Audio ID." % len(distinct),
      ["S.No", "Audio ID", "File name", "Hindi script (spoken exactly)", "Plays on", "Duration (s)", "Status"],
      {"A": 6, "B": 20, "C": 24, "D": 52, "E": 40, "F": 10, "G": 30})
for k, aid in enumerate(distinct, 1):
    pages = []
    for r in ROWS:
        if r[2] == aid and r[0] not in pages: pages.append(r[0])
    vals = [k, aid, aid + ".ogg", TEXT[aid], ", ".join(pages), dur(aid), status(aid)]
    for j, v in enumerate(vals, 1):
        c = ws.cell(4 + k, j, v); c.alignment = WRAP; c.border = BOX
        if j == 4: c.font = F_HI
# ---- Sound effects
ws = wb.create_sheet("Sound effects")
sheet(ws, "Sound effects and music (supplied - not recorded)", "Only the ones the lesson plays.",
      ["Audio ID", "File name", "What it is / where it plays", "Page(s)", "Duration (s)"],
      {"A": 20, "B": 22, "C": 80, "D": 12, "E": 11})
for k, (aid, what, pages) in enumerate(SFX, 1):
    for j, v in enumerate([aid, aid + ".ogg", what, pages, dur(aid)], 1):
        c = ws.cell(4 + k, j, v); c.alignment = WRAP; c.border = BOX
# ---- Recording spec
ws = wb.create_sheet("Recording spec")
ws["A1"] = "परी चली बाज़ार - recording spec"; ws["A1"].font = F_TITLE
spec = [
 ("Scope", "The whole lesson - start screen, 3 transitions, 8 pages (teaching, sort ₹10, sort ₹1, make ₹14 / ₹32 / ₹45, the shop game, celebration)."),
 ("Order", "Game order. On each page: what plays as it opens, then the hints in the order they can occur (Hint 1 / 2 / 3), then the closing line. A clip heard again later is listed once more with 'plays again'."),
 ("Clips", "%d rows, %d distinct recordings (see 'Recording list')." % (len(ROWS), len(distinct))),
 ("Folder", "assets/Audio/<Audio ID>.ogg - the file name is the Audio ID."),
 ("Current takes", "%d of %d are the SME's recordings (assets/Audio/voiceovers/<id>.wav, trimmed and levelled into assets/Audio/<id>.ogg by scripts/prepare_vo.py); the rest are Gemini TTS placeholders (voice Leda)." % (sum(os.path.isfile(os.path.join(SRC, i + ".wav")) for i in distinct), len(distinct))),
 ("Deliver", "One file per Audio ID, named exactly as the Audio ID (e.g. vo_landing.wav). WAV preferred, mono; any sample rate - the developer converts to the build's format (24 kHz mono, about -16 LUFS, peaks at most -1 dBFS)."),
 ("Script", "Record exactly the text in the 'Hindi script' column - it is the build's own spoken text. Amounts are written out as spoken (₹14 = «चौदह रुपये»); the screen shows ₹ and digits."),
 ("Drafts", "vo_t1_note and vo_t1_coin (page 1) were drafted by the developer - please finalise the wording before recording."),
 ("Register", "आप form throughout. Exceptions, as specified by the SME: the shop opening line «परी पहुँच गई है बाज़ार। आओ, चलो खरीदारी शुरू करें।» and the standard end-screen dialogue «बहुत बढ़िया, दोस्त! तुमने कमाल कर दिया!»."),
 ("Variants", "vo_hint_ones_1 ... _5 are one hint said for the amount still missing (₹1-₹5). Usually only _2 (page 5) and _5 (page 6) are heard; the others play only when the child has already placed some ₹1 coins."),
 ("Timing", "No fixed length - after delivery the developer rebuilds and re-measures what follows the voice: Swiftee's lip-sync on the transitions and the celebration, the typing of the transition titles, and the ₹10 / ₹1 / total highlights on pages 4-6."),
 ("EAR-CHECK", "vo_b45_prompt (the TTS needed a retry) - listen if the placeholder is kept."),
 ("Regenerate", "PYTHONUTF8=1 python scripts/gen_vo_manifest.py   (generated %s)" % datetime.date.today().isoformat()),
]
for k, (a, b) in enumerate(spec, 3):
    ws.cell(k, 1, a).font = Font(bold=True); c = ws.cell(k, 2, b); c.alignment = Alignment(wrap_text=True, vertical="top")
ws.column_dimensions["A"].width = 16; ws.column_dimensions["B"].width = 120
wb.save(OUT)
print("  OK  %s - %d rows, %d recordings, %d sound effects" % (os.path.basename(OUT), len(ROWS), len(distinct), len(SFX)))
